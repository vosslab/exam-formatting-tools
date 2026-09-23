"""Render answer choices and matching prompts in DOCX documents."""

import docx
import docx.shared

import ef_tools.docx_builder
import ef_tools.docx_images
import ef_tools.docx_table_builder
import ef_tools.layout
import ef_tools.text_utils


#============================================
def add_choice_content(para: object, choice: object, sizer: object = None,
		font_family: str = None) -> None:
	"""Add text and optional image content for one choice.

	Args:
		para: The paragraph to add content to.
		choice: Choice text or a structured text/image choice.
		sizer: Optional image sizing policy.
		font_family: Optional face override for text runs.
	"""
	choice_text = ef_tools.layout.choice_text(choice)
	if choice_text:
		ef_tools.docx_builder.add_rich_text_runs(
			para, choice_text, font_family=font_family)
	image_path = ef_tools.layout.choice_image(choice)
	if image_path:
		if choice_text:
			para.add_run().add_break()
		run = para.add_run()
		if sizer is None:
			run.add_picture(image_path)
			return
		kwargs = sizer.kwargs(image_path)
		ef_tools.docx_images.place_inline_picture(run, image_path, kwargs)


#============================================
def add_matching_prompt(doc: object, prompt: object, prefix: str,
		sizer: object = None, native_tables: bool = False) -> None:
	"""Add one `___ N.` matching prompt row; the prompt may carry an image.

	A plain string prompt renders as text (hard breaks become paragraphs).
	A dict prompt ({'text', 'image'}) renders its text, then its image
	inline on the same row so a drawing (DNA strip, pedigree) sits next to
	the blank. The sizer owns how big that image prints, including the
	wide-strip height rule for ordinary images.
	"""
	if not isinstance(prompt, dict):
		ef_tools.docx_builder.add_rich_text_paragraphs(
			doc, 'Matching Prompt', prompt, prefix=prefix)
		return
	if (native_tables and prompt.get('html_table')
			and ef_tools.docx_table_builder.supports_html_table(prompt['html_table'])):
		para = doc.add_paragraph()
		para.style = doc.styles['Matching Prompt']
		para.add_run(prefix)
		prompt_text = ef_tools.layout.choice_text(prompt)
		if prompt_text:
			ef_tools.docx_builder.add_rich_text_runs(para, prompt_text)
		para.paragraph_format.keep_with_next = True
		ef_tools.docx_table_builder.add_html_table(doc, prompt['html_table'])
		return
	para = doc.add_paragraph()
	para.style = doc.styles['Matching Prompt']
	para.add_run(prefix)
	add_choice_content(para, prompt, sizer=sizer)


#============================================
def _matching_prompt_can_share_row(prompt: object) -> bool:
	"""Return whether a plain-text prompt can share a two-column row."""
	return isinstance(prompt, str) and not ef_tools.text_utils.contains_hard_break(prompt)


#============================================
def _add_matching_prompt_text(para: object, prompt: str, prefix: str) -> None:
	"""Add one plain-text matching prompt to an existing paragraph."""
	para.add_run(prefix)
	ef_tools.docx_builder.add_rich_text_runs(para, prompt)


#============================================
def add_matching_prompts(doc: object, prompts: list, start_number: int,
		sizer: object = None, native_tables: bool = False) -> None:
	"""Add matching prompts, pairing simple text prompts on each row.

	Two plain-text prompts share a paragraph so the numbered blanks occupy two
	columns. Prompts with images, tables, or hard paragraph breaks retain the
	one-prompt-per-paragraph behavior because their content needs a full row.
	"""
	index = 0
	while index < len(prompts):
		prompt = prompts[index]
		next_prompt = prompts[index + 1] if index + 1 < len(prompts) else None
		if (_matching_prompt_can_share_row(prompt)
				and _matching_prompt_can_share_row(next_prompt)):
			para = doc.add_paragraph()
			para.style = doc.styles['Matching Prompt']
			_add_matching_prompt_text(
				para, prompt, f"___ {start_number + index}. ")
			para.add_run("\t")
			_add_matching_prompt_text(
				para, next_prompt, f"___ {start_number + index + 1}. ")
			para.paragraph_format.keep_with_next = index + 2 < len(prompts)
			index += 2
			continue
		add_matching_prompt(
			doc, prompt, f"___ {start_number + index}. ",
			sizer=sizer, native_tables=native_tables)
		# Hard-break prompts can emit more than one paragraph; only the final
		# paragraph needs to chain to the next prompt.
		if doc.paragraphs:
			doc.paragraphs[-1].paragraph_format.keep_with_next = (
				index + 1 < len(prompts))
		index += 1


# Wide table drawings such as DNA sequence strips are readable when stacked
# at page width, but become illegible when forced into the ordinary 4/5-column
# image-choice row. The threshold is supplied by styles/exam_styles.yaml.
# It is deliberately an aspect-ratio rule rather than a filename rule so the
# same layout works for any generator's strip drawing. The per-column width
# budgets live in ef_tools.docx_images.


#============================================
def _is_meaningful_alt(choice: object) -> bool:
	"""True when a choice's text adds information beyond a placeholder.

	Skips empty strings and the literal "image" so a row of placeholder
	alt text does not render as a redundant caption line. Used to decide
	whether an alt-text paragraph is emitted under an image-choice row.
	"""
	if not ef_tools.layout.choice_image(choice):
		return False
	text = (ef_tools.layout.choice_text(choice) or '').strip().lower()
	if not text or text == 'image':
		return False
	return True


#============================================
def _captioned_choices_need_stacking(choices: list) -> bool:
	"""Return whether a caption cannot fit one five-column choice slot.

	The threshold reuses the established five-column visible-width budget. A
	long caption in the shared tabbed caption row wraps from the page margin,
	so it can visually collide with the next choice; stacking gives each image
	and caption an unambiguous label relationship.
	"""
	return any(
		_is_meaningful_alt(choice)
		and ef_tools.text_utils.choice_visible_width(choice)
			> ef_tools.layout.DEFAULT_MAX_CHARS_5
		for choice in choices)


#============================================
def _add_captioned_image_choices_stacked(doc: docx.Document, choices: list,
		sizer: object) -> None:
	"""Render long-caption image choices as individually labeled blocks."""
	style_name = ef_tools.layout.choices_style_name(1)
	for index, choice in enumerate(choices):
		para = doc.add_paragraph()
		para.style = doc.styles[style_name]
		para.paragraph_format.space_after = docx.shared.Pt(0)
		letter = chr(ord('A') + index)
		prefix = para.add_run(f"({letter})")
		prefix.bold = True
		meaningful_alt = _is_meaningful_alt(choice)
		image_path = ef_tools.layout.choice_image(choice)
		if image_path:
			image_run = para.add_run()
			ef_tools.docx_images.place_inline_picture(
				image_run, image_path, sizer.kwargs(image_path))
		elif ef_tools.layout.choice_text(choice):
			choice_text = ef_tools.layout.choice_text(choice)
			para.add_run(' ')
			ef_tools.docx_builder.add_rich_text_runs(para, choice_text)
		if meaningful_alt:
			caption = doc.add_paragraph()
			caption.style = doc.styles[style_name]
			caption.paragraph_format.space_before = docx.shared.Pt(0)
			caption.paragraph_format.space_after = docx.shared.Pt(0)
			ef_tools.docx_builder.add_rich_text_runs(
				caption, ef_tools.layout.choice_text(choice))


#============================================
def _add_wide_image_choices_stacked(doc: docx.Document, choices: list,
		sizer: object) -> None:
	"""Render wide image choices one per row at a readable page width."""
	style_name = ef_tools.layout.choices_style_name(1)
	for index, choice in enumerate(choices):
		para = doc.add_paragraph()
		para.style = doc.styles[style_name]
		meaningful_alt = _is_meaningful_alt(choice)
		if index > 0:
			para.paragraph_format.space_before = docx.shared.Pt(0)
		if index < len(choices) - 1 or meaningful_alt:
			para.paragraph_format.space_after = docx.shared.Pt(0)
		letter = chr(ord('A') + index)
		prefix = para.add_run(f"({letter})")
		prefix.bold = True
		image_path = ef_tools.layout.choice_image(choice)
		if image_path:
			image_run = para.add_run()
			ef_tools.docx_images.place_inline_picture(
				image_run, image_path, sizer.kwargs(image_path))
		else:
			choice_text = ef_tools.layout.choice_text(choice)
			if choice_text:
				para.add_run(' ')
				ef_tools.docx_builder.add_rich_text_runs(para, choice_text)
		if meaningful_alt:
			caption = doc.add_paragraph()
			caption.style = doc.styles[style_name]
			caption.paragraph_format.space_before = docx.shared.Pt(0)
			caption.paragraph_format.space_after = docx.shared.Pt(0)
			caption_prefix = caption.add_run(f"({letter}) ")
			caption_prefix.bold = True
			ef_tools.docx_builder.add_rich_text_runs(
				caption, ef_tools.layout.choice_text(choice))


#============================================
def add_image_choices_tabbed(doc: docx.Document, choices: list,
		sizer: object, wide_sizer: object = None,
		wide_image_min_aspect: float = None) -> None:
	"""Add image-based choices in a horizontal tab-stop layout.

	Letter prefix and image render on the same line for each column.
	When at least one choice carries meaningful alt text (per
	`_is_meaningful_alt`), a second paragraph below holds the alt-text
	row. Tab stops are inherited from the Choices N paragraph style; no
	paragraph-level tab stops are added (style stops at the same column
	positions cause Word to see two close-but-not-equal stops per column
	and images drift off-grid). No docx tables are created.

	Args:
		doc: The Document to add the paragraph to.
		choices: List of structured choice dicts with text and/or image keys.
		sizer: Decides how big each image prints.
		wide_sizer: Sizer used when every choice is a wide strip and the row
			is stacked at page width instead.
		wide_image_min_aspect: Aspect-ratio threshold for the wide-strip path.
	"""
	if (wide_sizer is not None and wide_image_min_aspect is not None
			and choices
			and all(ef_tools.docx_images.is_wide_image_choice(
					choice, wide_image_min_aspect)
				for choice in choices)):
		_add_wide_image_choices_stacked(doc, choices, wide_sizer)
		return
	if choices and _captioned_choices_need_stacking(choices):
		_add_captioned_image_choices_stacked(
			doc, choices,
			sizer.narrowed(
				ef_tools.docx_images.CAPTIONED_IMAGE_CHOICE_MAX_WIDTH))
		return
	# clamp column count to the legal Choices 2..5 range so the
	# resolved style is always concrete (never the abstract Choice base)
	num_cols = len(choices)
	# A fixed-scale drawing cannot be squeezed into a column, so for drawings
	# the per-column budget is a fit test rather than a cap. A row too wide for
	# its columns stacks one per line, which keeps every drawing at the shared
	# scale instead of giving each one its own. A fitted image always fits,
	# because it is scaled into whatever column it gets.
	if choices and not sizer.fits_columns(choices, num_cols):
		_add_wide_image_choices_stacked(doc, choices, sizer)
		return
	style_columns = max(2, min(num_cols, 5))
	style_name = ef_tools.layout.choices_style_name(style_columns)
	# Narrow the sizer by the per-num_cols empirical budget. Column counts
	# beyond 5 (matching questions with many panels) fall back to the 5-col
	# value (the most conservative empirical cap).
	per_cols_cap = ef_tools.docx_images.IMAGE_CHOICE_MAX_WIDTH_BY_COLS.get(
		num_cols, ef_tools.docx_images.IMAGE_CHOICE_MAX_WIDTH_BY_COLS[5])
	column_sizer = sizer.narrowed(per_cols_cap)

	# Single combined paragraph: (A)<img> \t (B)<img> \t (C)<img> \t ...
	# Letter and image render side-by-side at each tab stop. Inherits
	# left_indent from the Choices N style so image-choice rows align
	# vertically with text-only Choices rows above/below them. Tab stops
	# in styles/exam_styles.yaml -> layout_tab_stops are pre-offset by
	# choice_indent so all inter-column gaps render evenly.
	combined_para = doc.add_paragraph()
	combined_para.style = doc.styles[style_name]
	combined_para.paragraph_format.space_after = docx.shared.Pt(0)
	# unify height across the row: a fitted sizer reports the smallest height
	# any image in the row would take, and every image is forced to it, which
	# keeps the row visually level when source aspect ratios differ. A drawing
	# sizer reports None, because unifying drawing heights would rescale each
	# one differently.
	row_height = column_sizer.row_height(choices)
	for index, choice in enumerate(choices):
		if index > 0:
			combined_para.add_run("\t")
		letter = chr(ord('A') + index)
		image_path = ef_tools.layout.choice_image(choice)
		# Image-bearing choices: bold "(A)" sits on the same line as the
		# image, no trailing space. Text-only choices: bold "(A) " plus
		# caption text on the same line, separator space.
		prefix_text = f"({letter})" if image_path else f"({letter}) "
		prefix = combined_para.add_run(prefix_text)
		prefix.bold = True
		if image_path:
			image_run = combined_para.add_run()
			if row_height is not None:
				kwargs = {'height': docx.shared.Inches(row_height)}
			else:
				kwargs = column_sizer.kwargs(image_path)
			# place_inline_picture also zeroes the inline image's
			# edge-distance margins so the image hugs the run, and applies a
			# signed offset to center it on the line's visual midline.
			ef_tools.docx_images.place_inline_picture(
				image_run, image_path, kwargs)
		else:
			choice_text = ef_tools.layout.choice_text(choice)
			if choice_text:
				ef_tools.docx_builder.add_rich_text_runs(
					combined_para, choice_text)

	# alt-text paragraph: rendered only when at least one image-bearing
	# choice carries non-trivial text. Skipping the row when every alt is
	# just "image" or empty avoids a redundant line of placeholder text.
	if any(_is_meaningful_alt(choice) for choice in choices):
		alt_para = doc.add_paragraph()
		alt_para.style = doc.styles[style_name]
		alt_para.paragraph_format.space_before = docx.shared.Pt(0)
		for index, choice in enumerate(choices):
			if index > 0:
				alt_para.add_run("\t")
			if ef_tools.layout.choice_image(choice):
				choice_text = ef_tools.layout.choice_text(choice)
				if choice_text:
					ef_tools.docx_builder.add_rich_text_runs(
						alt_para, choice_text)


#============================================
def add_choices_paragraph(doc: docx.Document, choices: list,
		tab_style: int, items_per_row: int, sizer: object = None,
		font_family: str = None) -> None:
	"""Add a tab-separated choices paragraph with bold letter prefixes.

	Creates a paragraph with (A) (B) (C) format, using tab characters
	to align columns. The tab_style selects tab stop positions (3, 4, 5)
	and items_per_row controls how many choices per line.

	Args:
		doc: The Document to add the paragraph to.
		choices: List of choice text strings or dicts with text/image keys.
		tab_style: Tab stop layout (3, 4, or 5) from CHOICES_TAB_STOPS.
		items_per_row: Number of choices per line (may differ from tab_style).
		sizer: Decides how big any choice image prints.
		font_family: Optional face override for choice text and letter labels.
	"""
	# any image in the list forces a vertical stack so the image and
	# its caption stay on one line per choice
	has_images = any(ef_tools.layout.choice_image(choice) for choice in choices)
	if has_images:
		items_per_row = 1
		tab_style = 1
	# select the appropriate style via the central resolver, which
	# always returns a concrete Choices N (never the bare Choice base)
	style_name = ef_tools.layout.choices_style_name(tab_style)
	# group choices into rows so each row becomes its own paragraph
	# (hard breaks). items_per_row<=1 means one choice per row.
	row_size = max(1, items_per_row)
	rows = [choices[start:start + row_size]
		for start in range(0, len(choices), row_size)]
	for row_index, row_choices in enumerate(rows):
		para = doc.add_paragraph()
		para.style = doc.styles[style_name]
		# rows after the first should sit flush against the prior row;
		# all rows except the last drop trailing space so the choices
		# group reads as one unit visually
		if row_index > 0:
			para.paragraph_format.space_before = docx.shared.Pt(0)
		if row_index < len(rows) - 1:
			para.paragraph_format.space_after = docx.shared.Pt(0)
		for col_index, choice in enumerate(row_choices):
			# global letter index across all rows
			letter = chr(ord('A') + row_index * row_size + col_index)
			# tab before this choice (except first in each row)
			if col_index > 0:
				para.add_run("\t")
			# bold letter prefix (only run-level override needed)
			bold_run = para.add_run(f"({letter}) ")
			bold_run.bold = True
			if font_family:
				bold_run.font.name = font_family
			# choice content inherits font size from style
			add_choice_content(
				para, choice, sizer=sizer, font_family=font_family)
