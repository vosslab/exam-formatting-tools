"""Convert simple HTML tables into native python-docx tables.

This is the optional backup path for bptools drawing tables. The normal
pipeline still keeps the rasterized PNG fallback; this module handles the
controlled table vocabulary emitted by the biology problem generators and
rejects nested or malformed layouts rather than guessing.
"""

# Standard Library
import html as html_lib

# Pip Modules
import docx
import docx.shared
import docx.enum.text
import docx.oxml
import docx.oxml.ns
import lxml.html

# Local Repo Modules
import ef_tools.bbq_html
import ef_tools.layout


class UnsupportedHtmlTable(ValueError):
	"""Raised when an HTML table is outside the native-table contract."""


#============================================
def has_native_table_choice(choice: object) -> bool:
	"""Return whether a choice carries one preserved HTML table."""
	return isinstance(choice, dict) and bool(choice.get('html_table'))


#============================================
def supports_html_table(table_html: str) -> bool:
	"""Return whether the conservative native-table parser accepts HTML."""
	try:
		_table_cells(table_html)
	except UnsupportedHtmlTable:
		return False
	return True


#============================================
def _style_properties(style: str) -> dict:
	"""Return lower-case CSS declarations from an inline style string."""
	properties = {}
	for declaration in style.split(';'):
		if ':' not in declaration:
			continue
		name, value = declaration.split(':', 1)
		properties[name.strip().lower()] = value.strip()
	return properties


#============================================
def _cell_background(cell: lxml.html.HtmlElement) -> str | None:
	"""Return a six-digit cell background color, if one is declared."""
	value = cell.get('bgcolor', '')
	properties = _style_properties(cell.get('style', ''))
	value = properties.get('background-color', value).strip()
	if value.startswith('#') and len(value) in (4, 7):
		if len(value) == 4:
			value = '#' + ''.join(character * 2 for character in value[1:])
		return value[1:].upper()
	return None


#============================================
def _span_value(cell: lxml.html.HtmlElement, name: str, default: int = 1) -> int:
	"""Read a positive integer rowspan/colspan attribute."""
	value = cell.get(name, str(default))
	try:
		parsed = int(value)
	except (TypeError, ValueError) as exc:
		raise UnsupportedHtmlTable(f"invalid {name}={value!r}") from exc
	if parsed < 1:
		raise UnsupportedHtmlTable(f"invalid {name}={value!r}")
	return parsed


#============================================
def _direct_rows(table: lxml.html.HtmlElement) -> list:
	"""Return table rows while respecting thead/tbody/tfoot boundaries."""
	rows = []
	for child in table:
		if child.tag == 'tr':
			rows.append(child)
		elif child.tag in ('thead', 'tbody', 'tfoot'):
			rows.extend(row for row in child if row.tag == 'tr')
		elif isinstance(child.tag, str) and child.tag not in ('colgroup', 'caption'):
			raise UnsupportedHtmlTable(f"unsupported table child <{child.tag}>")
	if not rows:
		raise UnsupportedHtmlTable('table has no rows')
	return rows


#============================================
def _reject_browser_layout(table: lxml.html.HtmlElement) -> None:
	"""Reject drawings whose meaning depends on browser positioning/CSS."""
	if table.get('role') == 'img':
		raise UnsupportedHtmlTable('role=img drawing requires raster fallback')
	if table.xpath('.//colgroup'):
		raise UnsupportedHtmlTable('colgroup layout requires raster fallback')
	if table.xpath('.//div'):
		raise UnsupportedHtmlTable('div-based layout requires raster fallback')
	for element in table.xpath('.//*[@style]'):
		properties = _style_properties(element.get('style', ''))
		if 'position' in properties or 'transform' in properties:
			raise UnsupportedHtmlTable('positioned layout requires raster fallback')


#============================================
def _cell_markup(cell: lxml.html.HtmlElement) -> str:
	"""Serialize only a cell's inline content, not its td wrapper."""
	parts = []
	if cell.text:
		parts.append(html_lib.escape(cell.text))
	for child in cell:
		parts.append(lxml.html.tostring(child, encoding='unicode', method='html'))
		if child.tail:
			parts.append(html_lib.escape(child.tail))
	return ''.join(parts)


#============================================
def _table_cells(table_html: str) -> tuple:
	"""Normalize an HTML table to a grid of logical cell records."""
	try:
		wrapper = lxml.html.fromstring(f'<div>{table_html}</div>')
	except (ValueError, TypeError) as exc:
		raise UnsupportedHtmlTable('table HTML could not be parsed') from exc
	tables = wrapper.xpath('.//table')
	if len(tables) != 1 or tables[0].xpath('.//table'):
		raise UnsupportedHtmlTable('nested tables are not supported')
	table = tables[0]
	_reject_browser_layout(table)
	rows = _direct_rows(table)
	occupied = [[None] * 0 for _ in rows]
	records = []
	max_columns = 0
	for row_index, row in enumerate(rows):
		column_index = 0
		for cell in row:
			if cell.tag not in ('td', 'th'):
				raise UnsupportedHtmlTable(f"unsupported row child <{cell.tag}>")
			while (column_index < len(occupied[row_index])
					and occupied[row_index][column_index] is not None):
				column_index += 1
			row_span = _span_value(cell, 'rowspan')
			column_span = _span_value(cell, 'colspan')
			if row_index + row_span > len(rows):
				raise UnsupportedHtmlTable('rowspan extends beyond the table')
			needed_columns = column_index + column_span
			if needed_columns > max_columns:
				max_columns = needed_columns
			for target_row in range(row_index, row_index + row_span):
				while len(occupied[target_row]) < needed_columns:
					occupied[target_row].append(None)
				for target_column in range(column_index, needed_columns):
					if occupied[target_row][target_column] is not None:
						raise UnsupportedHtmlTable('rowspan/colspan cells overlap')
					occupied[target_row][target_column] = (row_index, column_index)
			text = ef_tools.bbq_html.clean_inline_html(_cell_markup(cell))
			properties = _style_properties(cell.get('style', ''))
			records.append({
				'row': row_index,
				'column': column_index,
				'rowspan': row_span,
				'colspan': column_span,
				'text': text,
				'header': cell.tag == 'th',
				'background': _cell_background(cell),
				'alignment': properties.get('text-align'),
			})
			column_index = needed_columns
	return records, len(rows), max_columns


#============================================
def _shade_cell(cell: object, color: str) -> None:
	"""Set a Word cell's background fill."""
	shading = docx.oxml.OxmlElement('w:shd')
	shading.set(docx.oxml.ns.qn('w:fill'), color)
	shading.set(docx.oxml.ns.qn('w:val'), 'clear')
	cell._element.get_or_add_tcPr().append(shading)


#============================================
def _add_cell_text(cell: object, text: str, header: bool) -> None:
	"""Write rich text and hard-break paragraphs into one Word cell."""
	# Imported lazily to avoid a module cycle: docx_builder owns the rich-text
	# run emitter and imports this table backend at module load time.
	import ef_tools.docx_builder
	paragraphs = ef_tools.docx_builder._PARAGRAPH_BREAK_RE.split(text)
	paragraphs = [paragraph for paragraph in paragraphs if paragraph.strip()]
	if not paragraphs:
		paragraphs = ['']
	first = cell.paragraphs[0]
	for index, paragraph_text in enumerate(paragraphs):
		paragraph = first if index == 0 else cell.add_paragraph()
		paragraph.paragraph_format.space_after = docx.shared.Pt(0)
		ef_tools.docx_builder.add_rich_text_runs(paragraph, paragraph_text)
		if header:
			for run in paragraph.runs:
				run.font.bold = True


#============================================
def add_html_table(doc: object, table_html: str) -> object:
	"""Append one supported HTML table as a native Word table.

		Returns the created python-docx table. Nested and malformed tables raise
		UnsupportedHtmlTable so callers can retain the PNG fallback.
	"""
	records, row_count, column_count = _table_cells(table_html)
	table = doc.add_table(rows=row_count, cols=column_count)
	table.style = 'Table Grid'
	# Merge before writing text so each logical cell has one final anchor.
	for record in records:
		if record['rowspan'] > 1 or record['colspan'] > 1:
			start = table.cell(record['row'], record['column'])
			end = table.cell(
				record['row'] + record['rowspan'] - 1,
				record['column'] + record['colspan'] - 1)
			start.merge(end)
	for record in records:
		cell = table.cell(record['row'], record['column'])
		_add_cell_text(cell, record['text'], record['header'])
		if record['background']:
			_shade_cell(cell, record['background'])
		alignment = record['alignment']
		if alignment == 'center':
			for paragraph in cell.paragraphs:
				paragraph.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
		elif alignment == 'right':
			for paragraph in cell.paragraphs:
				paragraph.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.RIGHT
	return table


#============================================
def add_native_table_choices_stacked(doc: object, choices: list) -> None:
	"""Render table-based choices as one labeled native Word table each."""
	# Imported lazily to avoid a module cycle: the rich-text emitter lives in
	# docx_builder and imports this table backend at module load time.
	import ef_tools.docx_builder
	if not all(supports_html_table(choice['html_table']) for choice in choices):
		raise UnsupportedHtmlTable('one or more choice tables are unsupported')
	style_name = ef_tools.layout.choices_style_name(1)
	for index, choice in enumerate(choices):
		para = doc.add_paragraph()
		para.style = doc.styles[style_name]
		para.paragraph_format.space_after = docx.shared.Pt(0)
		label = para.add_run(f"({chr(ord('A') + index)})")
		label.bold = True
		choice_text = ef_tools.layout.choice_text(choice)
		if choice_text:
			para.add_run(" ")
			ef_tools.docx_builder.add_rich_text_runs(para, choice_text)
		add_html_table(doc, choice['html_table'])
