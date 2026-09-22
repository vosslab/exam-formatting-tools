"""Focused tests for the optional native HTML-table DOCX backend."""

import docx
import pytest

import ef_tools.docx_table_builder
import ef_tools.style_loader
import ef_tools.docx_builder


def _make_styled_doc() -> object:
	"""Create a document with the production exam styles."""
	doc = docx.Document()
	ef_tools.docx_builder.setup_styles(
		doc, ef_tools.style_loader.load_styles())
	return doc


def test_add_html_table_preserves_rich_text_colors_and_spans() -> None:
	"""Native tables retain cell text, color, shading, and colspan/rowspan."""
	table_html = (
		"<table><tr><th colspan='2' bgcolor='#E0E0E0'>A</th></tr>"
		"<tr><td rowspan='2'><span style='color: #004d00;'>C</span></td>"
		"<td>D</td></tr><tr><td>E</td></tr></table>"
	)
	doc = _make_styled_doc()
	table = ef_tools.docx_table_builder.add_html_table(doc, table_html)
	assert len(doc.tables) == 1
	assert len(table.rows) == 3
	assert table.cell(0, 0).text == 'A'
	assert table.cell(0, 1).text == 'A'
	assert table.cell(1, 0).text == 'C'
	assert table.cell(2, 1).text == 'E'
	assert table.cell(1, 0).paragraphs[0].runs[0].font.color.rgb == docx.shared.RGBColor(0, 77, 0)
	assert 'E0E0E0' in table.cell(0, 0)._tc.xml


def test_add_html_table_rejects_nested_tables() -> None:
	"""Unsupported nested layout raises so callers can keep the PNG fallback."""
	doc = _make_styled_doc()
	nested = '<table><tr><td><table><tr><td>x</td></tr></table></td></tr></table>'
	assert not ef_tools.docx_table_builder.supports_html_table(nested)
	with pytest.raises(ef_tools.docx_table_builder.UnsupportedHtmlTable):
		ef_tools.docx_table_builder.add_html_table(doc, nested)


def test_supports_html_table_rejects_browser_layout_drawings() -> None:
	"""Browser-positioned drawings retain their raster fallback."""
	role_image = (
		"<table role='img'><tr><td><div style='position:absolute'>x</div>"
		"</td></tr></table>")
	gel_layout = (
		"<table><colgroup width='5'></colgroup>"
		"<tr><td>x</td></tr></table>")
	assert not ef_tools.docx_table_builder.supports_html_table(role_image)
	assert not ef_tools.docx_table_builder.supports_html_table(gel_layout)
