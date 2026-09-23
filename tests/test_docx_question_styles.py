"""Tests for accessible question text and label formatting in DOCX output."""

import docx

import yaml_to_exam_docx


#============================================
def _exam_data() -> dict:
	"""Build a regular question followed by a matching question range."""
	result = {
		"title": "Question Style Test",
		"date": "2026-09-22",
		"sections": [
			{
				"questions": [
					{
						"statement": "Warm-up question.",
						"choices": ["A", "B"],
					},
					{
						"number": 12,
						"statement": (
							"Plain <code>ATGC</code> with <tt>GCTA</tt> and "
							"<b>explicit emphasis</b> and "
							"<strong><i><code>GCAT</code></i></strong>."
							"\nQuestion continuation."
						),
						"choices": ["A", "B"],
					},
					{
						"statement": "Match the terms.",
						"prompts_list": ["first", "second"],
						"choices_list": ["one", "two"],
					},
				],
			},
		],
	}
	return result


#============================================
def test_question_labels_are_boxed_and_question_text_uses_configured_emphasis(
		tmp_path: object) -> None:
	"""Question labels get a 2pt box; question prose uses accessible fonts."""
	output_path = tmp_path / "question_styles.docx"
	yaml_to_exam_docx.build_document(_exam_data(), str(output_path))
	doc = docx.Document(str(output_path))
	paragraphs = {para.text: para for para in doc.paragraphs}
	question_para = paragraphs[
		"12. Plain ATGC with GCTA and explicit emphasis and GCAT."]
	follow_para = paragraphs["Question continuation."]
	matching_para = paragraphs["Q13-14. Match the terms."]

	assert question_para.style.name == "Question Heading"
	assert follow_para.style.name == "Question Follow"
	assert matching_para.style.name == "Question Heading"
	question_heading_style = doc.styles["Question Heading"]
	question_follow_style = doc.styles["Question Follow"]
	assert question_heading_style.font.bold is False
	assert question_heading_style.font.italic is False
	assert question_follow_style.base_style == question_heading_style
	assert question_follow_style.font.bold is None
	assert question_follow_style.font.italic is None
	assert question_heading_style.paragraph_format.keep_with_next is True
	assert question_follow_style.paragraph_format.keep_with_next is True

	label_run = question_para.runs[0]
	assert label_run.text == "12."
	assert label_run.style.name == "Question Number"
	assert doc.styles["Question Number"].font.bold is False
	assert doc.styles["Question Number"].font.italic is False
	matching_label_run = matching_para.runs[0]
	assert matching_label_run.text == "Q13-14."
	assert matching_label_run.style.name == "Question Number"
	border = doc.styles["Question Number"].element.get_or_add_rPr().find(
		docx.oxml.ns.qn("w:bdr"))
	assert border is not None
	assert border.get(docx.oxml.ns.qn("w:val")) == "single"
	assert border.get(docx.oxml.ns.qn("w:sz")) == "16"
	assert border.get(docx.oxml.ns.qn("w:space")) == "2"

	for style_name in (
		"Normal", "Heading 1", "Exam Heading 2", "Chapter Heading", "Header",
		"Question Number",
	):
		assert doc.styles[style_name].font.name == "Atkinson Hyperlegible Next"
	code_runs = [run for run in question_para.runs if run.text in ("ATGC", "GCTA")]
	assert len(code_runs) == 2
	assert all(run.font.name == "Atkinson Hyperlegible Mono" for run in code_runs)
	assert all(run.font.bold is None for run in code_runs)
	assert all(run.font.italic is None for run in code_runs)
	assert all(style.name != "Exam Code" for style in doc.styles)
	explicit_code_run = next(
		run for run in question_para.runs if run.text == "GCAT")
	assert explicit_code_run.font.name == "Atkinson Hyperlegible Mono"
	assert explicit_code_run.bold is True
	assert explicit_code_run.italic is True
	explicit_bold_run = next(
		run for run in question_para.runs if run.text == "explicit emphasis")
	assert explicit_bold_run.bold is True

	prompt_para = next(para for para in doc.paragraphs if para.text.startswith("___ 13."))
	assert all(run.style.name != "Question Number" for run in prompt_para.runs)
