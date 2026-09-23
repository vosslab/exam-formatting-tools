"""Protect a synthetic task-CSV-to-YAML-to-DOCX workflow."""

import pathlib
import sys

import docx
import yaml

import bbq_tasks_to_exam_yaml
import ef_tools.bbq_parse
import ef_tools.bbq_tasks
import yaml_to_exam_docx


#============================================
class EmptyTableRenderer:
	"""Stand in for the renderer when the captured tasks contain no tables."""

	def __enter__(self) -> object:
		return self

	def __exit__(self, exc_type: object, exc_value: object,
			traceback: object) -> bool:
		return False


#============================================
def test_task_csv_aliases_reach_exam_yaml_and_docx(
		tmp_path: pathlib.Path, monkeypatch: object) -> None:
	"""YMATCH and YWHICH each become one rendered question from a fake bank."""
	problem_root = tmp_path / 'biology-problems' / 'problems'
	matching_script = problem_root / 'matching_sets' / 'yaml_match_to_bbq.py'
	which_script = problem_root / 'matching_sets' / 'yaml_which_one_mc_to_bbq.py'
	matching_script.parent.mkdir(parents=True)
	matching_script.write_text('# synthetic generator')
	which_script.write_text('# synthetic generator')

	task_csv = tmp_path / 'tasks.csv'
	task_csv.write_text(
		'subject,topic,script,flags,input,notes,choice_font\n'
		'genetics,genetic_disorders,YMATCH,,,,\n'
		'genetics,genetic_disorders,YWHICH,,,,IBM Plex Sans Condensed\n',
		encoding='utf-8')
	settings_path = tmp_path / 'settings.yml'
	settings_path.write_text(yaml.safe_dump({
		'paths': {'bp_root': str(problem_root)},
		'script_aliases': {
			'YMATCH': str(matching_script),
			'YWHICH': str(which_script),
		},
	}), encoding='utf-8')

	lines = {
		'yaml_match_to_bbq.py': (
			'MAT\t<p>Match the terms.</p>\tFirst prompt\tFirst answer'
			'\tSecond prompt\tSecond answer'),
		'yaml_which_one_mc_to_bbq.py': (
			'MC\t<p>Which statement is correct?</p>'
			'\tCorrect choice\tCorrect\tOther choice\tIncorrect'),
	}

	def fake_generate_question(task: dict, pythonpath: str,
			reject_fn: object) -> tuple:
		line = lines[pathlib.Path(task['script']).name]
		return ef_tools.bbq_parse.parse_bbq_line(line), ''

	class_renderer = bbq_tasks_to_exam_yaml.qti_package_maker.html_to_image.render_table
	monkeypatch.setattr(class_renderer, 'TableRenderer', EmptyTableRenderer)
	monkeypatch.setattr(
		ef_tools.bbq_tasks, 'generate_question', fake_generate_question)
	monkeypatch.setattr(
		ef_tools.bbq_parse.random, 'sample',
		lambda population, count: list(population)[:count])

	yaml_path = tmp_path / 'quiz.yml'
	monkeypatch.setattr(sys, 'argv', [
		'bbq_tasks_to_exam_yaml.py', '-i', str(task_csv), '-s',
		str(settings_path), '-o', str(yaml_path), '--quiz', '--title',
		'Synthetic task quiz'])
	bbq_tasks_to_exam_yaml.main()

	exam_data = yaml.safe_load(yaml_path.read_text(encoding='utf-8'))
	questions = exam_data['sections'][0]['questions']
	assert exam_data['title'] == 'Synthetic task quiz'
	assert exam_data['sections'][0]['chapter'] == 'Genetic disorders'
	assert len(questions) == 2
	assert questions[0]['prompts_list'] == [
		'First prompt', 'Second prompt']
	assert questions[1]['choices'] == ['Correct choice', 'Other choice']
	assert questions[1]['choice_font'] == 'IBM Plex Sans Condensed'
	answer_key = yaml_path.with_name('quiz-key.txt').read_text(encoding='utf-8')
	assert 'Q1-2. 1=A 2=B' in answer_key
	assert '3. A' in answer_key

	docx_path = tmp_path / 'quiz.docx'
	yaml_to_exam_docx.build_document(
		exam_data, str(docx_path), base_dir=str(tmp_path))
	document = docx.Document(docx_path)
	paragraph_text = [paragraph.text for paragraph in document.paragraphs]
	assert any('Match the terms.' in text for text in paragraph_text)
	assert any('First prompt' in text for text in paragraph_text)
	assert any('Which statement is correct?' in text for text in paragraph_text)
	assert any('Correct choice' in text for text in paragraph_text)
	choice_paragraph = next(
		paragraph for paragraph in document.paragraphs
		if 'Correct choice' in paragraph.text)
	assert any(
		run.font.name == 'IBM Plex Sans Condensed'
		for run in choice_paragraph.runs if 'Correct choice' in run.text)
	assert all('<p>' not in text for text in paragraph_text)
