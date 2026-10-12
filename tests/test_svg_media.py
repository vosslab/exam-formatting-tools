"""Protect generated SVG placement and media persistence without a browser."""

import zipfile

import docx
import pytest

import ef_tools.svg_media
import ef_tools.docx_images
import ef_tools.bbq_html
import ef_tools.bbq_parse
import docx_to_exam_yaml


def test_svg_preserves_physical_size_and_survives_docx_roundtrip(tmp_path: object) -> None:
	"""SVG drawings retain their source scale and are embedded without rasterization."""
	data = (b'<svg xmlns="http://www.w3.org/2000/svg" width="192px" height="96px" '
		b'viewBox="0 0 144 72"><path d="M0 0H144V72H0Z"/></svg>')
	image = tmp_path / 'drawing.svg'
	image.write_bytes(data)
	sizer = ef_tools.docx_images.DrawingSizer(0.5)
	assert sizer.width_inches(str(image)) == 1.0
	assert ef_tools.docx_images.rendered_height_inches(str(image), sizer.kwargs(str(image))) == 0.5
	document = docx.Document()
	ef_tools.docx_images.place_inline_picture(document.add_paragraph().add_run(),
		str(image), sizer.kwargs(str(image)))
	output = tmp_path / 'exam.docx'
	document.save(output)
	reopened = docx.Document(output)
	assert reopened.inline_shapes[0].width.inches == 1.0
	assert reopened.inline_shapes[0].height.inches == 0.5
	extracted = docx_to_exam_yaml.extract_images(reopened, str(tmp_path / 'extracted'))
	assert str(next(iter(extracted.values()))).endswith('.svg')
	with zipfile.ZipFile(output) as archive:
		assets = [name for name in archive.namelist() if name.startswith('word/media/')]
		assert len(assets) == 1
		assert archive.read(assets[0]) == data


def test_svg_captures_are_saved_with_correct_extension(tmp_path: object) -> None:
	"""Both table and standalone media keep SVG bytes rather than mislabeled PNGs."""
	data = (b'<svg xmlns="http://www.w3.org/2000/svg" width="96px" height="48px" '
		b'viewBox="0 0 72 36"><path d="M0 0H72V36Z"/></svg>')
	media_dir = tmp_path / 'exam_files'
	paths = ef_tools.bbq_html.write_table_images(
		['<table data-ef-image="1"><tr><td>A</td></tr></table>'],
		None, str(media_dir), 'question', {'1': data})
	assert paths == ['exam_files/question_table_1.svg']
	reference = ef_tools.bbq_html.RENDERED_IMAGE_PREFIX + '1'
	path = ef_tools.bbq_parse._save_rendered_image(reference, {'1': data},
		str(media_dir), str(tmp_path), 'equation.png', {})
	assert path == 'exam_files/equation.svg'
	assert (tmp_path / path).read_bytes() == data


def test_svg_fit_uses_the_same_aspect_ratio_as_raster_images(tmp_path: object) -> None:
	"""A vector drawing fits its box without stretching either axis."""
	image = tmp_path / 'equation.svg'
	image.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="192px" height="96px"/>')
	sizer = ef_tools.docx_images.FitSizer(3, max_height=1)
	assert sizer.width_inches(str(image)) == pytest.approx(2)
