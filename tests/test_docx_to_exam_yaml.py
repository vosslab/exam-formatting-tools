"""DOCX import keeps statement text and image paragraphs in order."""

import docx

import docx_to_exam_yaml


#============================================
def test_parse_docx_questions_emits_ordered_statement_blocks() -> None:
	"""A paragraph image remains between the question and its continuation."""
	document = docx.Document()
	document.add_paragraph('1. Which molecule is shown?')
	document.add_paragraph('')
	document.add_paragraph('The label appears below the figure.')
	questions = docx_to_exam_yaml.parse_docx_questions(
		document, {1: 'figure.png'})
	assert questions[0]['statement'] == [
		{'text': '1. Which molecule is shown?'},
		{'image': 'figure.png'},
		{'text': 'The label appears below the figure.'},
	]


#============================================
def test_parse_docx_questions_keeps_tables_in_statement_order() -> None:
	"""An imported Word table stays between its surrounding paragraphs."""
	document = docx.Document()
	document.add_paragraph('1. What happens in the pathway?')
	document.add_paragraph('The table comes next.')
	table = document.add_table(rows=2, cols=1)
	table.cell(0, 0).text = 'Step'
	table.cell(1, 0).text = 'Product'
	document.add_paragraph('The second paragraph follows the table.')
	questions = docx_to_exam_yaml.parse_docx_questions(document, {})
	statement = questions[0]['statement']
	assert statement[0] == {'text': '1. What happens in the pathway?'}
	assert statement[1] == {'text': 'The table comes next.'}
	assert statement[2] == {
		'table': {'columns': ['Step'], 'rows': [['Product']]}}
	assert statement[3] == {'text': 'The second paragraph follows the table.'}
