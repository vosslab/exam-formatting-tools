"""Protect the synthetic HTML-to-YAML-to-DOCX delivery path."""

import sys

import docx
import PIL.Image
import yaml

import html_to_exam_yaml
import yaml_to_exam_docx


#============================================
def test_html_media_color_and_order_reach_docx(tmp_path: object,
		monkeypatch: object) -> None:
	"""Generated YAML resolves statement media and preserves authored order."""
	source_dir = tmp_path / 'source'
	media_dir = source_dir / 'files'
	output_dir = tmp_path / 'output'
	media_dir.mkdir(parents=True)
	output_dir.mkdir()
	for image_name, color in (('first.png', 'purple'), ('second.png', 'green')):
		PIL.Image.new('RGB', (12, 8), color).save(media_dir / image_name)

	html_path = source_dir / 'question.html'
	html_path.write_text(
		'<html><head><title>Pipeline</title></head><body>'
		'<div class="takeQuestionDiv"><li>'
		'<p>Before <span style="color: 800080;">Metabolite X</span>.</p>'
		'<img class="cleaned-statement-media" src="files/first.png">'
		'<p>Between the images.</p>'
		'<img class="cleaned-statement-media" src="files/second.png">'
		'<p>After both images.</p>'
		'<div><label><input type="radio"><span>A. Yes</span></label>'
		'<label><input type="radio"><span>B. No</span></label></div>'
		'</li></div></body></html>',
		encoding='utf-8')
	yaml_path = output_dir / 'quiz.yml'
	monkeypatch.setattr(sys, 'argv', [
		'html_to_exam_yaml.py', '-i', str(html_path), '-o', str(yaml_path),
		'-t', 'Synthetic pipeline', '--date', '2026-09-22'])
	html_to_exam_yaml.main()
	exam_data = yaml.safe_load(yaml_path.read_text(encoding='utf-8'))
	statement = exam_data['sections'][0]['questions'][0]['statement']
	assert [next(iter(block)) for block in statement] == [
		'text', 'image', 'text', 'image', 'text']
	assert statement[0]['text'] == (
		'Before <span style="color: #800080;">Metabolite X</span>.')
	assert statement[1]['image'] == '../source/files/first.png'
	assert statement[3]['image'] == '../source/files/second.png'

	docx_path = output_dir / 'quiz.docx'
	yaml_to_exam_docx.build_document(
		exam_data, str(docx_path), base_dir=str(output_dir))
	document = docx.Document(docx_path)
	question_content = [
		paragraph for paragraph in document.paragraphs
		if ('Before' in paragraph.text or 'Between' in paragraph.text
			or 'After both' in paragraph.text
			or paragraph._element.xpath('.//w:drawing'))]
	assert len(question_content) == 5
	assert 'Before' in question_content[0].text
	assert question_content[1]._element.xpath('.//w:drawing')
	assert 'Between' in question_content[2].text
	assert question_content[3]._element.xpath('.//w:drawing')
	assert 'After both' in question_content[4].text
	colored_run = next(
		run for paragraph in document.paragraphs for run in paragraph.runs
		if run.text == 'Metabolite X')
	assert str(colored_run.font.color.rgb) == '800080'
